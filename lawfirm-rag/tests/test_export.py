"""Offline tests for draft export (DOCX / Markdown / text).

No live DB, no network. A .docx is a zip; we assert the bytes are a valid zip,
re-open the document to confirm the headings/endnotes/footer landed, and check
the text formats contain the same structure.
"""
from __future__ import annotations

import zipfile
from io import BytesIO

from docx import Document

from drafting import export


DRAFT = {
    "title": "Acme / Beta mutual NDA",
    "doc_type": "nda",
    "status": "draft",
    "current_version": 2,
    "sections": [
        {
            "order_index": 0,
            "heading": "Parties",
            "content": "This Agreement is between Acme Ltd and Beta Ltd [1].\n\nBoth are Kenyan companies.",
            "sources": [{"n": 1, "title": "Acme precedent", "section_heading": "Parties", "file_id": "f1"}],
            "status": "generated",
        },
        {
            "order_index": 1,
            "heading": "Governing Law",
            "content": "",
            "sources": [],
            "status": "empty",
        },
    ],
}


def _docx_paragraphs(data: bytes) -> list[str]:
    document = Document(BytesIO(data))
    return [p.text for p in document.paragraphs]


def test_build_docx_is_a_valid_zip():
    data = export.build_docx(DRAFT)
    assert data[:2] == b"PK"
    assert zipfile.is_zipfile(BytesIO(data))


def test_build_docx_contains_headings_and_title():
    paragraphs = _docx_paragraphs(export.build_docx(DRAFT))
    joined = "\n".join(paragraphs)
    assert "Acme / Beta mutual NDA" in joined
    assert "Parties" in joined
    assert "Governing Law" in joined


def test_build_docx_gathers_endnotes_and_footer():
    data = export.build_docx(DRAFT)
    paragraphs = _docx_paragraphs(data)
    joined = "\n".join(paragraphs)
    assert "Endnotes — sources" in joined
    assert "[1] Acme precedent" in joined
    # The standing disclaimer is written into the section footer.
    document = Document(BytesIO(data))
    footer_text = "\n".join(p.text for p in document.sections[0].footer.paragraphs)
    assert "Not legal advice" in footer_text


def test_build_docx_handles_empty_draft():
    data = export.build_docx({"title": "Empty", "sections": []})
    assert zipfile.is_zipfile(BytesIO(data))
    assert "no sections" in "\n".join(_docx_paragraphs(data))


def test_build_markdown_and_text_structure():
    md = export.build_markdown(DRAFT)
    assert "# Acme / Beta mutual NDA" in md
    assert "## Parties" in md
    assert "## Endnotes — sources" in md
    assert export.DRAFT_DISCLAIMER in md

    txt = export.build_text(DRAFT)
    assert "PARTIES" in txt
    assert "ENDNOTES - SOURCES" in txt
    assert export.DRAFT_DISCLAIMER in txt
