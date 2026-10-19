"""Page/paragraph provenance for citations.

The old ``metadata.page_number`` was ``chunk_index // chunks_per_page`` — a
guess, 0 for every DOCX/text file, and never the page the text came from.
These tests lock in the replacement: real page/paragraph spans derived from
``ingest.extract`` and mapped onto chunks by ``ingest.embed``, then surfaced in
the sources list, the prompt context, the viewer and the exports.
"""
from __future__ import annotations

import io

from langchain_core.documents import Document

from agents.rag_agent import (
    _build_context_string,
    _location_from_meta,
    _locator_label,
    _quote_from,
    _sources_from_docs,
)
from drafting.export import _source_line
from ingest.embed import (
    TextChunk,
    _apply_provenance,
    _attach_offsets,
    chunk_text,
    chunk_text_auto,
)
from ingest.extract import Segment, clean_text, extract_document, extract_text
from routers.documents import _chunk_locator_label


# ---------------------------------------------------------------------------
# Extraction keeps page/paragraph structure
# ---------------------------------------------------------------------------

def test_extract_plain_text_numbers_paragraphs():
    extracted = extract_document(b"First paragraph.\n\nSecond paragraph.\n\nThird.", "text/plain")

    assert extracted.locator_kind == "paragraph"
    assert [s.paragraph for s in extracted.segments] == [1, 2, 3]
    assert extracted.text == "First paragraph.\n\nSecond paragraph.\n\nThird."
    # Offsets must actually address the text.
    for segment in extracted.segments:
        assert extracted.text[segment.start:segment.end].strip()


def test_extract_docx_numbers_paragraphs():
    from docx import Document as Docx

    document = Docx()
    document.add_paragraph("Clause 1. The parties agree.")
    document.add_paragraph("")  # empty paragraphs are not numbered
    document.add_paragraph("Clause 2. The term is five years.")
    buffer = io.BytesIO()
    document.save(buffer)

    extracted = extract_document(
        buffer.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    assert extracted.locator_kind == "paragraph"
    assert [s.paragraph for s in extracted.segments] == [1, 2]
    assert "Clause 1" in extracted.text and "Clause 2" in extracted.text
    # DOCX has no intrinsic pages.
    assert extracted.total_pages == 1


def test_extract_text_wrapper_is_unchanged():
    text, pages = extract_text(b"Only paragraph.", "text/plain")
    assert text == "Only paragraph."
    assert pages == 1


def test_clean_text_still_public():
    assert clean_text("Hello\n\n\n\nWorld") == "Hello\n\nWorld"


def _build_pdf(pages: list[str]) -> bytes:
    """Build a small, valid PDF with one text line per page.

    Hand-rolled because no PDF-*writing* library is a dependency; pypdf is only
    used to read. This keeps the page-provenance path covered against real PDF
    bytes rather than a synthetic stand-in.
    """
    objects: list[tuple[int, bytes]] = []
    page_ids: list[int] = []
    content_ids: list[int] = []
    nxt = 3
    for _ in range(len(pages)):
        page_ids.append(nxt)
        nxt += 1
        content_ids.append(nxt)
        nxt += 1
    font_id = nxt

    def obj(num: int, body: bytes | str) -> None:
        objects.append((num, body if isinstance(body, bytes) else body.encode()))

    obj(1, b"<< /Type /Catalog /Pages 2 0 R >>")
    kids = " ".join(f"{pid} 0 R" for pid in page_ids)
    obj(2, f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>")
    for index, text in enumerate(pages):
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
        obj(
            page_ids[index],
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Contents {content_ids[index]} 0 R "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> >>",
        )
        obj(
            content_ids[index],
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        )
    obj(font_id, b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    out = bytearray(b"%PDF-1.4\n")
    offsets: dict[int, int] = {}
    for num, body in sorted(objects):
        offsets[num] = len(out)
        out += f"{num} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_offset = len(out)
    size = max(offsets) + 1
    out += f"xref\n0 {size}\n".encode() + b"0000000000 65535 f \n"
    for num in range(1, size):
        out += f"{offsets[num]:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {size} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n"
    ).encode()
    return bytes(out)


def test_pdf_extraction_reports_real_pages():
    raw = _build_pdf(
        ["Alpha page one content", "Beta page two content", "Gamma page three content"]
    )

    extracted = extract_document(raw, "application/pdf")

    assert extracted.locator_kind == "page"
    assert extracted.total_pages == 3
    assert [s.page for s in extracted.segments] == [1, 2, 3]
    # Each segment must address its own page's words, not merely overlap the text.
    assert "Alpha" in extracted.text[extracted.segments[0].start:extracted.segments[0].end]
    assert "Beta" in extracted.text[extracted.segments[1].start:extracted.segments[1].end]
    assert "Gamma" in extracted.text[extracted.segments[2].start:extracted.segments[2].end]


def test_pdf_blank_page_does_not_shift_real_page_numbers():
    raw = _build_pdf(["Alpha first", " ", "Gamma third"])

    extracted = extract_document(raw, "application/pdf")

    assert extracted.total_pages == 3
    # The blank page produces no segment, but the surviving segments keep the
    # real page numbers (1 and 3) instead of being renumbered 1 and 2.
    assert [s.page for s in extracted.segments] == [1, 3]
    assert "Alpha" in extracted.text[extracted.segments[0].start:extracted.segments[0].end]
    assert "Gamma" in extracted.text[extracted.segments[1].start:extracted.segments[1].end]


def test_pdf_chunk_reports_the_page_it_came_from():
    raw = _build_pdf(
        ["Alpha page one content", "Beta page two content", "Gamma page three content"]
    )
    extracted = extract_document(raw, "application/pdf")

    chunks = chunk_text(
        extracted.text,
        chunk_size=8,
        chunk_overlap=0,
        segments=extracted.segments,
        locator_kind=extracted.locator_kind,
    )

    assert [(c.page_start, c.page_end) for c in chunks] == [(1, 1), (2, 2), (3, 3)]
    assert chunks[1].text.strip() == "Beta page two content"


def test_pdf_chunk_straddling_pages_reports_a_range():
    raw = _build_pdf(
        ["Alpha page one content", "Beta page two content", "Gamma page three content"]
    )
    extracted = extract_document(raw, "application/pdf")

    chunks = chunk_text(
        extracted.text,
        chunk_size=10,
        chunk_overlap=0,
        segments=extracted.segments,
        locator_kind=extracted.locator_kind,
    )

    # The first chunk holds all of page 1 and part of page 2: report the range.
    assert (chunks[0].page_start, chunks[0].page_end) == (1, 2)


# ---------------------------------------------------------------------------
# Chunking maps back onto the segments
# ---------------------------------------------------------------------------

def test_chunk_reports_page_range_when_straddling_pages():
    text = ("alpha " * 20).strip() + "\n\n" + ("beta " * 20).strip()
    midpoint = text.index("\n\n") + 1
    segments = (
        Segment(start=0, end=midpoint, page=3),
        Segment(start=midpoint + 1, end=len(text), page=4),
    )

    chunks = chunk_text(text, chunk_size=5000, chunk_overlap=0, segments=segments, locator_kind="page")

    assert len(chunks) == 1
    assert (chunks[0].page_start, chunks[0].page_end) == (3, 4)
    assert chunks[0].locator_kind == "page"
    assert chunks[0].char_start == 0


def test_chunks_carry_monotonic_paragraphs():
    text = "\n\n".join(f"Paragraph number {i} with a little more text." for i in range(1, 13))
    extracted = extract_document(text.encode(), "text/plain")

    # One chunk small enough to hold a few paragraphs, so several chunks result.
    chunks = chunk_text(
        text, chunk_size=25, chunk_overlap=0, segments=extracted.segments, locator_kind="paragraph"
    )

    assert len(chunks) > 1
    starts = [c.paragraph_start for c in chunks]
    assert all(s is not None for s in starts)
    assert starts == sorted(starts)
    assert all(c.locator_kind == "paragraph" for c in chunks)


def test_single_chunk_spans_the_whole_document():
    text = "\n\n".join(f"Paragraph number {i} with a little more text." for i in range(1, 13))
    extracted = extract_document(text.encode(), "text/plain")

    chunks = chunk_text_auto(text, None, segments=extracted.segments, locator_kind=extracted.locator_kind)[0]

    assert len(chunks) == 1
    assert (chunks[0].paragraph_start, chunks[0].paragraph_end) == (1, 12)


def test_chunk_without_segments_has_no_locator():
    chunks = chunk_text("Just some text with no provenance.")
    assert chunks[0].page_start is None
    assert chunks[0].paragraph_start is None
    assert chunks[0].locator_kind == "none"


def test_attach_offsets_recovers_positions_for_agentic_chunks():
    text = "alpha beta gamma delta"
    chunks = [TextChunk("beta", "", 0), TextChunk("gamma", "", 1)]

    located = _attach_offsets(chunks, text)

    assert (located[0].char_start, located[0].char_end) == (6, 10)
    assert (located[1].char_start, located[1].char_end) == (11, 16)


def test_unlocatable_chunk_gets_no_locator():
    # A chunk the agentic splitter cannot be found in the source must not be
    # pinned to page 1 — an absent page is honest, a wrong one is not.
    located = _attach_offsets([TextChunk("not in the text", "", 0)], "alpha beta")
    assert located[0].char_start == -1

    segments = (Segment(start=0, end=10, page=1),)
    with_provenance = chunk_text(
        "alpha beta", segments=segments, locator_kind="page"
    )
    assert with_provenance[0].page_start == 1  # sanity: real chunks do locate

    applied = _apply_provenance(located, segments, "page")
    assert applied[0].page_start is None
    assert applied[0].page_end is None


# ---------------------------------------------------------------------------
# Labels and source assembly
# ---------------------------------------------------------------------------

def test_locator_label_variants():
    assert _locator_label(12, 12, None, None) == "p. 12"
    assert _locator_label(12, 13, None, None) == "pp. 12–13"
    assert _locator_label(None, None, 4, 4) == "¶ 4"
    assert _locator_label(None, None, 4, 6) == "¶¶ 4–6"
    assert _locator_label(None, None, None, None) == ""
    # Pages win when a document somehow carries both.
    assert _locator_label(2, 2, 9, 9) == "p. 2"


def test_location_from_meta_ignores_legacy_zero_page_number():
    # Pre-provenance DOCX/TXT chunks had page_number = 0: unknown, not page 0.
    assert _location_from_meta({"page_number": 0}) is None
    assert _location_from_meta({"file_title": "x"}) is None
    assert _location_from_meta({"page_start": 7, "page_end": 8})["label"] == "pp. 7–8"


def test_quote_from_truncates_on_a_word_boundary():
    assert _quote_from("  a   b  ") == "a b"
    long = " ".join(["word"] * 200)
    quote = _quote_from(long)
    assert quote.endswith("…")
    assert len(quote) <= 321
    assert not quote[:-1].endswith(" ")


def _doc(**meta):
    return Document(page_content=meta.pop("content", "some retrieved text"), metadata=meta)


def test_sources_aggregate_locations_and_quote_for_one_document():
    docs = [
        _doc(file_id="f1", file_title="Lease.pdf", content="Rent is 280,000 [first].", page_start=12, page_end=13, chunk_index=3, locator_kind="page"),
        _doc(file_id="f1", file_title="Lease.pdf", content="Service charge clause.", page_start=14, page_end=14, chunk_index=4, locator_kind="page"),
        _doc(file_id="f2", file_title="Judgment.pdf", content="The appeal is dismissed.", page_start=1, page_end=1, chunk_index=0, locator_kind="page"),
    ]

    sources = _sources_from_docs(docs)

    assert [s["n"] for s in sources] == [1, 2]
    lease = sources[0]
    assert lease["file_title"] == "Lease.pdf"
    assert (lease["page_start"], lease["page_end"]) == (12, 14)
    assert lease["label"] == "pp. 12–14"
    assert lease["quote"] == "Rent is 280,000 [first]."
    assert len(lease["locations"]) == 2
    assert sources[1]["label"] == "p. 1"


def test_sources_without_provenance_show_no_page():
    sources = _sources_from_docs([_doc(file_id="upload_abc_file.pdf", file_title="Scan.pdf", page_number=0)])
    assert sources[0]["label"] == ""
    assert sources[0]["locations"] == []
    assert sources[0]["page_number"] == 0


def test_context_string_includes_the_location():
    docs = [
        _doc(file_id="f1", file_title="Lease.pdf", page_start=12, page_end=13, chunk_index=0, locator_kind="page"),
    ]
    sources = _sources_from_docs(docs)

    context = _build_context_string(docs, sources)

    assert "[1] Lease" in context
    assert "| pp. 12–13" in context


def test_context_string_omits_location_when_unknown():
    docs = [_doc(file_id="f1", file_title="Scan.pdf", page_number=0)]
    context = _build_context_string(docs, _sources_from_docs(docs))
    assert "| p." not in context and "| ¶" not in context


# ---------------------------------------------------------------------------
# Exports and the in-app viewer
# ---------------------------------------------------------------------------

def test_export_source_line_uses_the_label():
    assert _source_line({"title": "Lease.pdf", "label": "pp. 12–13"}) == "Lease.pdf — pp. 12–13"
    assert _source_line({"title": "Notes.docx", "label": "¶ 4"}) == "Notes.docx — ¶ 4"


def test_export_source_line_falls_back_for_legacy_sections():
    # Sections persisted before the label field existed still render a locator.
    assert _source_line({"title": "Lease.pdf", "page_start": 12, "page_end": 13}) == "Lease.pdf — pp. 12–13"
    assert _source_line({"title": "Lease.pdf", "section_heading": "3. RENT"}) == "Lease.pdf — 3. RENT"


def test_viewer_chunk_locator_label():
    assert _chunk_locator_label({"page_start": 4}) == "p. 4"
    assert _chunk_locator_label({"page_start": 4, "page_end": 6}) == "pp. 4–6"
    assert _chunk_locator_label({"paragraph_start": 2, "paragraph_end": 3}) == "¶¶ 2–3"
    assert _chunk_locator_label({"page_number": 0}) == ""
    assert _chunk_locator_label({}) == ""
